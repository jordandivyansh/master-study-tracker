import openpyxl, json, sys, datetime
src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
def d(v):
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime('%Y-%m-%d')
    return v
def num(v):
    return v if isinstance(v,(int,float)) and not isinstance(v,bool) else None
def dump(name, obj):
    json.dump(obj, open(f'{out}/{name}.json','w'), indent=2, ensure_ascii=False)
    print(name, len(obj) if isinstance(obj,list) else 'obj')

S = wb['Start']; v = lambda r: S.cell(r,2).value
settings = {
  'dailyCap': v(8), 'dsaWeekly': v(9), 'buffer': v(10), 'gymDaily': v(11),
  'trackerStart': d(v(12)), 'trackerEnd': d(v(13)),
  'exams': [
    {'key':'ssc','name':'SSC CGL Tier I','date':d(v(14)),'prepBudgetH':v(15),'note':S.cell(15,4).value},
    {'key':'cat','name':'CAT','date':d(v(16)),'prepBudgetH':v(17),'note':S.cell(17,4).value},
    {'key':'uppsc','name':'UPPSC Prelims','date':d(v(18)),'prepBudgetH':None,'note':None},
    {'key':'apfc','name':'UPSC EPFO APFC','date':d(v(19)),'prepBudgetH':120,'note':'APFC course budget for selected gaps, PYQs, mocks and review.'},
  ],
  'pivotExam':'apfc',
  'weight': {'start':v(20),'target':v(21),'goalDate':d(v(22))},
  'cava': {'newPerWeek':v(23),'hoursPerLecture':v(24)},
  'durations': {
    'CSE':{'attempt':v(25),'review':v(26)},
    'UPPSC':{'attempt':v(27),'review':v(28)},
    'CAT':{'attempt':v(29),'review':v(30)},
    'APFC':{'attempt':v(31),'review':v(32)},
    'SSC':{'attempt':v(33),'review':v(34)},
  },
}
dump('settings', settings)

W = wb['Weekly']; keys = ['ssc','cat','uppsc','apfc','cava','pubad','mains','books','other']
weeks=[]
for r in range(8,31):
    row=[c.value for c in W[r]]
    weeks.append({'start':d(row[0]),'end':d(row[1]),'phase':row[2],
      'quotas':{k: (num(row[9+i]) or 0) for i,k in enumerate(keys)},'notes':row[26]})
dump('weeks', weeks)

B = wb['CSE Backlog']; backlog=[]
for i,r in enumerate(range(8,19)):
    row=[c.value for c in B[r]]
    e=row[4]
    hpu = settings['cava']['hoursPerLecture'] if isinstance(e,str) else num(e)
    backlog.append({'order':i+1,'workstream':row[0],'unit':row[1],'totalUnits':num(row[2]),'completedUnits':num(row[3]),
      'hoursPerUnit':hpu,'progressOverride': num(row[6]) if not isinstance(row[6],str) else None,
      'status':row[7],'deadline':d(row[8]),'notes':row[9],'source':row[10]})
dump('backlog', backlog)

T = wb['Tests']; tests=[]
for r in range(8,67):
    row=[c.value for c in T[r]]
    if not row[0]: continue
    tests.append({'code':row[0],'exam':row[1],'paper':row[2],'type':row[3],'release':d(row[4]),
      'targetAttempt':d(row[5]),'targetReview':d(row[6]),'attemptStatus':row[7],'reviewStatus':row[8],
      'actualAttempt':d(row[9]),'actualReview':d(row[10]),'maxScore':num(row[15]),'score':num(row[16]),'notes':row[18]})
dump('tests', tests)

A = wb['APFC Course']; mods=[]
for r in range(8,37):
    row=[c.value for c in A[r]]
    mods.append({'num':row[0],'module':row[1],'items':row[2],'platformComplete':row[3],'attachmentBadges':row[5],
      'priority':row[6],'action':row[7],'studyStatus':row[8],'verifiedRuntimeH':num(row[9]),'selectedStudyH':num(row[10]),
      'hoursDone':num(row[11]),'notes':row[13]})
dump('apfcModules', mods)
areas=[]
for r in range(50,59):
    row=[c.value for c in A[r]]
    areas.append({'num':row[0],'area':row[1],'firstCoverage':d(row[2]),'questionsReviewed':num(row[3]),
      'errorsCorrected':d(row[4]),'recallChecked':d(row[5]),'action':row[7]})
dump('syllabusAreas', areas)

C = wb['CGB Entries']; entries=[]
for r in range(8, C.max_row+1):
    row=[c.value for c in C[r]]
    if not row[0] or not str(row[0]).startswith('CGB-'): continue
    entries.append({'code':row[0],'module':row[1],'title':row[2],'type':row[3],'platformStatus':row[4],
      'displayedDuration':row[5],'observedRuntimeH':num(row[6]),'studyMode':row[7],'selectedStudyH':num(row[8]),
      'hoursDone':num(row[9]),'reviewStatus':row[11],'notes':row[12]})
dump('courseEntries', entries)

D = wb['Daily']; logs=[]
dk = ['ssc','cat','uppsc','apfc','cava','pubad','mains','books','other']
for r in range(8,164):
    row=[c.value for c in D[r]]
    if not row[0]: continue
    hrs={k:num(row[2+i]) for i,k in enumerate(dk)}
    if any(x is not None for x in list(hrs.values())+[num(row[11]),num(row[15]),num(row[16])]) or row[17]:
        logs.append({'date':d(row[0]),'hours':hrs,'dsa':num(row[11]),'gym':num(row[15]),'weight':num(row[16]),'notes':row[17]})
dump('dailyLogs', logs)
